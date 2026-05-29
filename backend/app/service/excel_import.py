"""Leitura e validação de planilhas .xlsx para importação de despesas.

Mapeia o cabeçalho de forma tolerante (sem acento, case-insensitive, aceitando
rótulos como "Data de Vencimento" ou "Valor (R$)"), normaliza valores em formato
BR e devolve os schemas válidos junto com os erros linha a linha.
"""
import io
import re
import unicodedata
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook

from app.schema.despesa import DespesaSchema, ImportRowError


CANONICAL_FIELDS = ("nome", "tipo", "valor", "vencimento", "status", "descricao")
REQUIRED_FIELDS = ("nome", "tipo", "valor", "vencimento")
TEMPLATE_HEADERS = ["nome", "tipo", "valor", "vencimento", "status", "descricao"]

MAX_NOME = 100
MAX_TIPO = 100
MAX_DESCRICAO = 20
DEFAULT_DESCRICAO = "PARCELA ÚNICA"

_STATUS_MAP = {
    "p": "P",
    "q": "Q",
    "pendente": "P",
    "aberta": "P",
    "aberto": "P",
    "quitada": "Q",
    "quitado": "Q",
    "paga": "Q",
    "pago": "Q",
}


def _strip_accents_lower(value: Any) -> str:
    s = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    return s.strip().lower()


def _map_headers(header_row) -> dict[str, int]:
    """Associa cada campo canônico ao índice da coluna no cabeçalho."""
    mapping: dict[str, int] = {}
    for idx, raw in enumerate(header_row):
        if raw is None:
            continue
        norm = _strip_accents_lower(raw)
        if not norm:
            continue
        for field in CANONICAL_FIELDS:
            if field in mapping:
                continue
            if norm == field or re.search(rf"\b{field}\b", norm):
                mapping[field] = idx
                break
    return mapping


def _parse_valor(raw: Any) -> float:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        raise ValueError("valor é obrigatório")
    if isinstance(raw, bool):
        raise ValueError(f"valor inválido: '{raw}'")
    if isinstance(raw, (int, float)):
        valor = float(raw)
    else:
        s = str(raw).strip().replace("R$", "").replace(" ", "")
        # Formato BR: ponto de milhar, vírgula decimal.
        s = s.replace(".", "").replace(",", ".")
        try:
            valor = float(s)
        except ValueError:
            raise ValueError(f"valor inválido: '{raw}'")
    if valor <= 0:
        raise ValueError("valor deve ser maior que zero")
    return valor


def _parse_vencimento(raw: Any) -> date:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        raise ValueError("vencimento é obrigatório")
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip().split(" ")[0]
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(s.split("T")[0])
    except ValueError:
        raise ValueError(f"vencimento inválido: '{raw}' (use dd/mm/aaaa)")


def _parse_status(raw: Any) -> str:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return "P"  # default: Pendente
    key = _strip_accents_lower(raw)
    if key in _STATUS_MAP:
        return _STATUS_MAP[key]
    raise ValueError(f"status inválido: '{raw}' (use Pendente/Quitada ou P/Q)")


def _parse_texto_obrigatorio(raw: Any, campo: str, maxlen: int) -> str:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        raise ValueError(f"{campo} é obrigatório")
    texto = str(raw).strip()
    if len(texto) > maxlen:
        raise ValueError(f"{campo} excede {maxlen} caracteres")
    return texto


def _is_empty_row(values) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in values)


def parse_despesas_xlsx(
    content: bytes,
) -> tuple[list[DespesaSchema], list[ImportRowError], int]:
    """Lê o .xlsx e devolve (despesas válidas, erros por linha, total de linhas).

    Levanta ValueError para problemas que invalidam o arquivo inteiro
    (arquivo ilegível, vazio ou sem as colunas obrigatórias).
    """
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Não foi possível ler o arquivo. Envie um .xlsx válido.")

    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)

        try:
            header = next(rows)
        except StopIteration:
            raise ValueError("A planilha está vazia.")

        mapping = _map_headers(header)
        faltando = [f for f in REQUIRED_FIELDS if f not in mapping]
        if faltando:
            raise ValueError(
                "Colunas obrigatórias ausentes: "
                + ", ".join(faltando)
                + ". Esperado: "
                + ", ".join(TEMPLATE_HEADERS)
            )

        def cell(values, field):
            idx = mapping.get(field)
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        validos: list[DespesaSchema] = []
        erros: list[ImportRowError] = []
        total = 0

        for linha, values in enumerate(rows, start=2):  # linha 1 = cabeçalho
            if _is_empty_row(values):
                continue
            total += 1
            try:
                nome = _parse_texto_obrigatorio(cell(values, "nome"), "nome", MAX_NOME)
                tipo = _parse_texto_obrigatorio(cell(values, "tipo"), "tipo", MAX_TIPO)
                valor = _parse_valor(cell(values, "valor"))
                vencimento = _parse_vencimento(cell(values, "vencimento"))
                status = _parse_status(cell(values, "status"))

                descricao = DEFAULT_DESCRICAO
                descricao_raw = cell(values, "descricao")
                if descricao_raw is not None and str(descricao_raw).strip():
                    descricao = str(descricao_raw).strip()
                    if len(descricao) > MAX_DESCRICAO:
                        raise ValueError(f"descricao excede {MAX_DESCRICAO} caracteres")

                validos.append(
                    DespesaSchema(
                        nome=nome,
                        tipo=tipo,
                        valor=valor,
                        vencimento=vencimento,
                        status=status,
                        descricao=descricao,
                    )
                )
            except ValueError as error:
                erros.append(ImportRowError(linha=linha, erro=str(error)))
            except Exception as error:  # erros de validação do pydantic, etc.
                erros.append(ImportRowError(linha=linha, erro=str(error).split("\n")[0]))

        return validos, erros, total
    finally:
        wb.close()


def build_template_xlsx() -> bytes:
    """Gera um modelo .xlsx com o cabeçalho esperado e linhas de exemplo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Despesas"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["Conta de Luz", "BOLETO", "150,75", "31/12/2026", "Pendente", "PARCELA ÚNICA"])
    ws.append(["Notebook", "CARTAO", "3.000,00", "10/01/2027", "Quitada", "PARCELA 1/10"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
